#!/usr/bin/env python3
"""把手册 Markdown 草稿构建为操作手册 DOCX（python-docx 后端，无 OfficeCLI 依赖）。

支持章节标题、正文、表格、图片（本地相对路径）、```text 等宽块（架构图/代码）。
版式与代码材料一致：A4，上下边距 2.0cm，左 2.5cm、右 2.0cm；
页眉左侧“软件全称 版本号”，右侧“第 <页码> 页”（PAGE 字段，Word/WPS/LibreOffice 均可重排）。

用法：
  python3 build_manual_docx.py --md 草稿/操作手册.md \
      --software-name "<软件全称>" --version V1.0 \
      --out 预览资料/<软件全称>_操作手册.docx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

ASCII_BLOCK_MAX_COLS = 78


def set_run_font(run, size_pt, bold=False, mono=False, ea="SimSun"):
    ascii_font = "Courier New" if mono else "Times New Roman"
    if mono:
        ea = "SimSun"
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), ea)


def add_header(document: Document, title: str) -> None:
    section = document.sections[0]
    header = section.header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.text = ""
    usable = section.page_width - section.left_margin - section.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(title + "\t")
    set_run_font(run, 9)
    run2 = p.add_run("第 ")
    set_run_font(run2, 9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    frun = OxmlElement("w:r")
    ftext = OxmlElement("w:t")
    ftext.text = "1"
    frun.append(ftext)
    fld.append(frun)
    p._p.append(fld)
    run3 = p.add_run(" 页")
    set_run_font(run3, 9)


def add_body_paragraph(document: Document, text: str, size=10.5):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    for seg, bold in split_bold(text):
        run = p.add_run(seg)
        set_run_font(run, size, bold=bold)
    return p


def split_bold(text: str):
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    return [(p.strip("*"), "**" in p and len(p) > 4) for p in parts if p]


def add_ascii_block(document: Document, lines: list[str]):
    for ln in lines:
        p = document.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "240")
        spacing.set(qn("w:lineRule"), "exact")
        pPr.append(spacing)
        run = p.add_run(ln if ln.strip() else " ")
        set_run_font(run, 7.5, mono=True)


def add_image(document: Document, path: Path, alt: str):
    if not path.exists():
        add_body_paragraph(document, f"【界面图 {path.name} 缺失：{alt}】")
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(14.5))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    crun = cap.add_run(f"图 {alt}")
    set_run_font(crun, 9, bold=True)


def add_table(document: Document, rows: list[list[str]]):
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, cells in enumerate(rows):
        for c, cell_text in enumerate(cells):
            if c >= len(table.rows[r].cells):
                continue
            cell = table.rows[r].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(cell_text.replace("`", ""))
            set_run_font(run, 8.5, bold=(r == 0))


def build(md_path: Path, out_path: Path, software_name: str, version: str) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    add_header(document, f"{software_name} {version}")

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            i += 1
            add_ascii_block(document, block)
            continue
        m_img = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", s)
        if m_img:
            add_image(document, (md_path.parent / m_img.group(2)).resolve(), m_img.group(1))
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            add_table(document, rows)
            continue
        if s.startswith("# "):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(s[2:].strip())
            set_run_font(run, 16, bold=True)
            i += 1
            continue
        if s.startswith("## "):
            p = document.add_paragraph()
            run = p.add_run(s[3:].strip())
            set_run_font(run, 14, bold=True)
            p.paragraph_format.space_before = Pt(10)
            i += 1
            continue
        if s.startswith("### "):
            p = document.add_paragraph()
            run = p.add_run(s[4:].strip())
            set_run_font(run, 12, bold=True)
            p.paragraph_format.space_before = Pt(6)
            i += 1
            continue
        if s.startswith("> "):
            i += 1
            continue
        add_body_paragraph(document, s)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--md", type=Path, required=True)
    ap.add_argument("--software-name", required=True)
    ap.add_argument("--version", default="V1.0")
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()
    build(args.md, args.out, args.software_name, args.version)
    print(f"OK {args.out}")


if __name__ == "__main__":
    main()
